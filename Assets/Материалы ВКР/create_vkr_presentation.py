from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("ВКР_Шаповалов_презентация_защита.pptx")


slides = [
    {
        "title": "Реализация группового управления ботами в компьютерной игре",
        "subtitle": "ВКР, 09.03.02 Информационные системы и технологии\nШаповалов Михаил Андреевич, ИТси-422\nРуководитель: доцент, к.т.н. Дмитриев Н. В.\nУрГУПС, Екатеринбург, 2026",
        "visual": "Dungeon Seeker\n2D top-down RPG\nгрупповое управление NPC",
    },
    {
        "title": "Цель и задачи работы",
        "bullets": [
            "Цель: разработать 2D RPG Dungeon Seeker и интегрировать интеллектуальную систему поведения NPC.",
            "1. Сравнить игровые движки и фреймворки для сложной ИИ-логики.",
            "2. Разработать архитектуру игрового приложения.",
            "3. Реализовать ИИ NPC с классическими методами и машинным обучением.",
            "4. Провести функциональное и нагрузочное тестирование.",
            "5. Оценить перспективы расширения игры.",
            "6. Подготовить документацию, экономику и БЖД.",
        ],
        "visual": "Цель -> 6 задач -> готовый игровой прототип",
    },
    {
        "title": "Задача 1. Выбор инструментов и методов",
        "bullets": [
            "Сравнены Unity, Godot и Unreal Engine по 2D-разработке, ИИ, навигации, ML и экосистеме.",
            "Выбран Unity 6: компонентная модель, 2D-инструменты, Unity Behavior, ML-Agents, NavMesh Plus.",
            "Для NPC обоснован гибридный подход: классический каркас + обучаемая микрозадача позиционирования.",
        ],
        "visual": "Сравнительная схема:\nUnity 6 = 2D + Behavior + ML-Agents + NavMesh",
    },
    {
        "title": "Задача 2. Архитектура приложения",
        "bullets": [
            "Сформирована компонентная архитектура игрового продукта.",
            "Базовые системы: управление, оружие, способности, инвентарь, экипировка, враги, группа, навигация, UI.",
            "Данные баланса вынесены в ScriptableObject, что упрощает настройку без изменения кода.",
        ],
        "visual": "Слои:\nUI и контент\nГеймплейные системы\nИИ NPC\nДанные ScriptableObject",
    },
    {
        "title": "Задача 3. Интеллектуальная система NPC",
        "bullets": [
            "Союзник: Unity Behavior организует действия, Utility AI выбирает цель и способность.",
            "ML-Agents управляет боевым позиционированием: дистанция, отступление, направление и интенсивность стрейфа.",
            "Противники реализованы через вычислительно дешевый и предсказуемый конечный автомат.",
        ],
        "visual": "Unity Behavior -> Utility AI -> ML-Agents -> NavMesh\nHFSM для массовых противников",
    },
    {
        "title": "Задача 3. Обучение ML-компонента",
        "bullets": [
            "Создана сцена ML_Training_Arena с эпизодами до 90 секунд и постепенным усложнением волн.",
            "Агент AllyPosition использует 17 наблюдений и 4 непрерывных действия.",
            "Обучение выполнено алгоритмом PPO до 1 000 054 шагов, результат экспортирован в ONNX.",
        ],
        "visual": "ML_Training_Arena:\nнаблюдения -> PPO -> политика AllyPosition.onnx -> Inference в игре",
    },
    {
        "title": "Задача 4. Тестирование и результаты",
        "bullets": [
            "Проведено сравнение эвристического и обученного режимов на 30 эпизодах максимальной сложности.",
            "Успешность выросла с 90% до 100%; смерти агента снизились с 3 до 0.",
            "Средняя награда выросла с 4,01 до 5,49, стандартное отклонение снизилось с 5,34 до 1,28.",
        ],
        "visual": "Таблица/диаграмма:\n90% -> 100%\n4,01 -> 5,49\n5,34 -> 1,28",
    },
    {
        "title": "Задача 5. Перспективы развития",
        "bullets": [
            "Процедурная генерация уровней для повышения реиграбельности и разнообразия тренировочных ситуаций.",
            "Расширение контента: новые классы, способности, предметы, враги и профили поведения.",
            "Перспектива многопользовательского режима с серверной инфраструктурой.",
        ],
        "visual": "Дорожная карта:\nконтент -> процедурная генерация -> мультиплеер",
    },
    {
        "title": "Задача 6. Безопасность жизнедеятельности",
        "bullets": [
            "В разделе БЖД рассмотрены риски серверной инфраструктуры: электротравмы, пожарная опасность, шум, микроклимат и статические перегрузки.",
            "Предложены меры: УЗО, защитное зануление, газовое пожаротушение, вынос рабочего места оператора и поддержание микроклимата 20-22 °C.",
        ],
        "visual": "Схема серверной:\nэлектробезопасность + пожарная защита + микроклимат",
        "speech": "Для безопасной эксплуатации перспективной серверной инфраструктуры предложены меры электробезопасности, пожарной защиты, снижения шума и поддержания нормального микроклимата.",
    },
    {
        "title": "Задача 6. Экономическое обоснование",
        "bullets": [
            "Трудоемкость проекта: 704 человеко-часа.",
            "Общие инвестиции: 816 352,57 руб.",
            "Расчетный срок окупаемости: 6 месяцев; чистый дисконтированный доход: 84 212 руб.",
        ],
        "visual": "Структура затрат:\nтруд и отчисления 49,6%\nмаркетинг 36,7%\nпрочие затраты",
        "speech": "Экономический расчет показал, что при инвестициях 816 352,57 рубля проект окупается на шестой месяц и имеет положительный чистый дисконтированный доход 84 212 рубля.",
    },
    {
        "title": "Заключение",
        "bullets": [
            "Цель достигнута: разработан игровой прототип Dungeon Seeker с интеллектуальной системой NPC.",
            "Решены все задачи: выбран стек, разработана архитектура, реализованы базовые механики и гибридный ИИ.",
            "Обученная ONNX-модель повысила устойчивость поведения союзника по сравнению с эвристикой.",
            "Проведены тестирование, оценка перспектив, анализ БЖД и экономическое обоснование.",
        ],
        "visual": "Все задачи -> подтвержденный результат -> дальнейшее развитие",
    },
]


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title(slide, text):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.35), Inches(12.2), Inches(0.75))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(245, 248, 255)


def add_bullets(slide, bullets):
    box = slide.shapes.add_textbox(Inches(0.65), Inches(1.35), Inches(7.3), Inches(5.35))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(235, 240, 247)
        p.space_after = Pt(8)


def add_visual(slide, text):
    shape = slide.shapes.add_shape(1, Inches(8.35), Inches(1.45), Inches(4.25), Inches(4.95))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(31, 46, 75)
    shape.line.color.rgb = RGBColor(96, 138, 209)
    shape.line.width = Pt(2)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(21)
    p.font.bold = True
    p.font.color.rgb = RGBColor(245, 248, 255)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for idx, data in enumerate(slides):
        slide = prs.slides.add_slide(blank)
        set_bg(slide, RGBColor(13, 24, 42))
        if idx == 0:
            title = slide.shapes.add_textbox(Inches(0.75), Inches(0.85), Inches(11.8), Inches(1.25))
            p = title.text_frame.paragraphs[0]
            p.text = data["title"]
            p.alignment = PP_ALIGN.CENTER
            p.font.size = Pt(34)
            p.font.bold = True
            p.font.color.rgb = RGBColor(245, 248, 255)

            subtitle = slide.shapes.add_textbox(Inches(1.25), Inches(2.25), Inches(10.8), Inches(1.45))
            p = subtitle.text_frame.paragraphs[0]
            p.text = data["subtitle"]
            p.alignment = PP_ALIGN.CENTER
            p.font.size = Pt(18)
            p.font.color.rgb = RGBColor(220, 230, 245)

            add_visual(slide, data["visual"])
        else:
            add_title(slide, data["title"])
            add_bullets(slide, data["bullets"])
            add_visual(slide, data["visual"])

        footer = slide.shapes.add_textbox(Inches(0.55), Inches(7.02), Inches(12.2), Inches(0.25))
        p = footer.text_frame.paragraphs[0]
        p.text = f"{idx + 1} / {len(slides)}"
        p.alignment = PP_ALIGN.RIGHT
        p.font.size = Pt(10)
        p.font.color.rgb = RGBColor(145, 162, 190)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
